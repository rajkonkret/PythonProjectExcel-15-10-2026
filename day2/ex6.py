import openpyxl

filename = 'video2.xlsx'
wb = openpyxl.load_workbook(filename)
# ws = wb.active
ws = wb['vgsales']

ws['P1'] = "Avaregaes Sales"
ws['P2'] = '=AVERAGE(K2:K16328)'
# =ŚREDNIA(K2:K16328)

wb.save(filename)
wb.close()

ws['Q1'] = "Number of populated cells"
ws['Q2'] = '=COUNTA(E2:E16328)'
# =ILE.NIEPUSTYCH(E2:E16328)

wb.save(filename)
wb.close()