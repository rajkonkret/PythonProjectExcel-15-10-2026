import openpyxl

wb = openpyxl.load_workbook('videogamesales.xlsx')
ws = wb.active

print(wb)
print(ws)
# <openpyxl.workbook.workbook.Workbook object at 0x0000027ED959A3C0>
# <Worksheet "vgsales">

ws = wb['vgsales']

print("Total number of rows:", ws.max_row)
print("Total number of columns:", ws.max_column)
# Total number of rows: 16328
# Total number of columns: 10

print("Value in cell A1 is:", ws['A1'].value)
# Value in cell A1 is: Rank

# list comprehensions
print([i for i in range(5)])
var = [i for i in range(5)]

values = [ws.cell(row=1, column=i).value for i in range(1, ws.max_column + 1)]
print(values)
# ['Rank', 'Name', 'Platform', 'Year', 'Genre', 'Publisher', 'NA_Sales', 'EU_Sales', 'JP_Sales', 'Other_Sales']

data = [ws.cell(row=i, column=2).value for i in range(2, 12)]
print(data)

# my_list = []
my_list = list()

for value in ws.iter_rows(
        min_row=1,
        max_row=11,
        min_col=1,
        max_col=6,
        values_only=True
):
    my_list.append(value)

print(my_list)

from pprint import pprint

pprint(my_list)

for ele1, ele2, ele3, ele4, ele5, ele6 in my_list:
    print(f"{ele1:<8}{ele2:>35}{ele3:^10}{ele4:<10}{ele5:<15}{ele6:<15}")

# dodanie nowej kolumny
ws["K1"] = "Sum of Sales"

wb.save('video2.xlsx')
wb.close()