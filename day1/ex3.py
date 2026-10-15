import openpyxl

filename = 'video2.xlsx'
wb = openpyxl.load_workbook(filename)
ws = wb.active

print(wb)
print(ws)

row_position = 2
col_position = 7

total_sales = (
        (ws.cell(row=row_position, column=col_position).value)
        + (ws.cell(row=row_position, column=col_position + 1).value)
        + (ws.cell(row=row_position, column=col_position + 2).value)
        + (ws.cell(row=row_position, column=col_position + 3).value)
)

ws.cell(row=2, column=11).value = total_sales

wb.save(filename)
wb.close()
